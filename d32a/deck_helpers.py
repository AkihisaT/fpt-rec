# -*- coding: utf-8 -*-
# pptx helpers (from the 003/036 decks): textbox, title, footer, picture, table. FOOT must be defined by the caller.
# Build the FPT #003 summary deck (16:9, Yu Gothic) from the bundle folder.
import os, re
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

FONT = "游ゴシック"
NAVY = RGBColor(0x1F, 0x38, 0x64); ACC = RGBColor(0x2F, 0x55, 0x97); GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xE9, 0xEE, 0xF6); WHITE = RGBColor(0xFF, 0xFF, 0xFF); BLACK = RGBColor(0x20, 0x20, 0x20)
SW, SH = 13.333, 7.5


def _font(run, size, bold=False, color=BLACK, italic=False):
    f = run.font; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color; f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None: e = etree.SubElement(rPr, qn(tag))
        e.set("typeface", FONT)

def add_runs(p, text, size, color=BLACK, bold=False):
    """**bold** segments inside text; '\n' becomes a line break (<a:br/>) inside the paragraph."""
    if "\n" in text:
        for k, line in enumerate(text.split("\n")):
            if k: p.add_line_break()
            add_runs(p, line, size, color, bold)
        return
    SUP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part: continue
        b = part.startswith("**") and part.endswith("**")
        seg = part[2:-2] if b else part
        for piece in re.split(r"(⁻[⁰¹²³⁴⁵⁶⁷⁸⁹]+)", seg):      # '⁻' is missing in Yu Gothic -> real superscript
            if not piece: continue
            r = p.add_run()
            if piece.startswith("⁻"):
                r.text = "-" + piece[1:].translate(SUP); _font(r, size, bold or b, color)
                r.font._rPr.set("baseline", "30000")
            else:
                r.text = piece; _font(r, size, bold or b, color)

def _bullet(p, level, char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.22 + 0.28 * level)))); pPr.set("indent", str(-int(Inches(0.2))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", char if level == 0 else "–")

def textbox(slide, x, y, w, h, items, size=15, color=BLACK, bullets=True, space=5, anchor=MSO_ANCHOR.TOP, align=None, line=1.08):
    """items: list of str or (level, str) or ('h', str) for a small heading line."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for it in items:
        lvl, txt = (0, it) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.line_spacing = line
        if lvl == "h":
            add_runs(p, txt, size, ACC, True); p.space_before = Pt(space + 2); continue
        add_runs(p, txt, size - (2 if lvl == 1 else 0), color)
        p.space_before = Pt(space)
        if align: p.alignment = align
        if bullets: _bullet(p, lvl)
    return tb

def title(slide, text, sub=None):
    textbox(slide, 0.45, 0.28, 12.4, 0.62, [text], size=26, color=NAVY, bullets=False, anchor=MSO_ANCHOR.MIDDLE)
    for p in slide.shapes[-1].text_frame.paragraphs:
        for r in p.runs: r.font.bold = True
    ln = slide.shapes.add_connector(1, Inches(0.5), Inches(0.95), Inches(SW - 0.5), Inches(0.95))
    ln.line.color.rgb = ACC; ln.line.width = Pt(1.5)
    if sub:
        textbox(slide, 0.5, 0.98, 12.3, 0.4, [sub], size=13, color=GREY, bullets=False)

def footer(slide, n):
    textbox(slide, 0.45, SH - 0.38, 9, 0.3, [FOOT], size=9, color=GREY, bullets=False)
    textbox(slide, SW - 1.2, SH - 0.38, 0.75, 0.3, [str(n)], size=9, color=GREY, bullets=False, align=PP_ALIGN.RIGHT)


# ---------- pictures ----------
def picture(slide, path, x, y, w=None, h=None):
    """Place image inside box (x, y, w, h) keeping aspect; returns (left, top, width, height) in inches."""
    iw, ih = Image.open(path).size
    if w and h:
        s = min(w / iw, h / ih); pw, ph = iw * s, ih * s
        x += (w - pw) / 2
    elif w: pw, ph = w, w * ih / iw
    else: ph, pw = h, h * iw / ih
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(pw), Inches(ph))
    return x, y, pw, ph

def split_comparison(path, tag):
    """Split a comparison figure into the image grid and the bottom metric row at the widest white band (60-92 % height)."""
    im = Image.open(path).convert("RGB"); a = np.asarray(im).astype(int)
    white = (a.min(axis=2) > 245).all(axis=1)
    H = a.shape[0]; lo, hi = int(0.60 * H), int(0.92 * H)
    best, run, start = (0, 0), 0, lo
    for r in range(lo, hi):
        if white[r]:
            if run == 0: start = r
            run += 1
            if run > best[1] - best[0]: best = (start, r + 1)
        else: run = 0
    cut = (best[0] + best[1]) // 2
    os.makedirs("/tmp/deckcrops", exist_ok=True)
    top, bot = f"/tmp/deckcrops/{tag}_grid.png", f"/tmp/deckcrops/{tag}_metrics.png"
    im.crop((0, 0, im.width, cut)).save(top); im.crop((0, cut, im.width, H)).save(bot)
    return top, bot, (best[1] - best[0]), cut / H


# ---------- tables ----------
def table(slide, rows, x, y, w, col_w=None, size=11, header=True, row_h=0.3, first_col_bold=True):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None: style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"      # 'No Style, Table Grid'
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w): tbl.columns[j].width = Inches(w * cw / tot)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, txt in enumerate(row):
            c = tbl.cell(i, j); c.margin_left = c.margin_right = Inches(0.05); c.margin_top = c.margin_bottom = Inches(0.025)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.line_spacing = 1.0
            hdr = header and i == 0
            add_runs(p, txt, size, WHITE if hdr else BLACK, bold=hdr or (first_col_bold and j == 0))
            c.fill.solid(); c.fill.fore_color.rgb = ACC if hdr else (LIGHT if i % 2 == 0 else WHITE)
    return shp
